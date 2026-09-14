"""Права на вкладку «Рационы Claude» и карточка блюда в редакторе рациона.

Ограниченный редактор — читатель с галочкой «Может вести рационы Claude»:
правит только свои группы и рационы плюс группы, отмеченные shared_editing
(это группы ДИАБЕТ). Чужое остаётся только для чтения, остальные вкладки —
как у обычного читателя.
"""

from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse

from .models import (
    ClaudeRation, ClaudeRationGroup, ClaudeRationSlot,
    MealTime, Product, UserRights,
)
from .roles import READERS_GROUP, can_edit_claude_group, can_edit_claude_ration


class ClaudeEditorRightsTests(TestCase):
    def setUp(self):
        readers, _ = Group.objects.get_or_create(name=READERS_GROUP)

        self.limited = User.objects.create_user("roman", password="x")
        self.limited.groups.add(readers)
        UserRights.objects.create(user=self.limited, can_edit_claude_rations=True)

        self.plain_reader = User.objects.create_user("saule", password="x")
        self.plain_reader.groups.add(readers)

        self.editor = User.objects.create_user("adil", password="x")  # не читатель

        self.diabet = ClaudeRationGroup.objects.create(
            name="ДИАБЕТ 1800", shared_editing=True,
        )
        self.foreign = ClaudeRationGroup.objects.create(name="Рационы 1200 (тест)")
        self.own = ClaudeRationGroup.objects.create(
            name="Моя группа", created_by=self.limited,
        )

        self.foreign_ration = ClaudeRation.objects.create(
            group=self.foreign, name="Чужой рацион", kcal_category=1200,
        )
        self.own_ration = ClaudeRation.objects.create(
            group=self.own, name="Свой рацион", kcal_category=1200,
            created_by=self.limited,
        )
        self.diabet_ration = ClaudeRation.objects.create(
            group=self.diabet, name="Диабет рацион", kcal_category=1800,
        )

    # ── правила доступа ──────────────────────────────────────────────────────

    def test_limited_editor_scope(self):
        self.assertTrue(can_edit_claude_group(self.limited, self.diabet))
        self.assertTrue(can_edit_claude_group(self.limited, self.own))
        self.assertFalse(can_edit_claude_group(self.limited, self.foreign))

        self.assertTrue(can_edit_claude_ration(self.limited, self.own_ration))
        self.assertTrue(can_edit_claude_ration(self.limited, self.diabet_ration))
        self.assertFalse(can_edit_claude_ration(self.limited, self.foreign_ration))

    def test_plain_reader_has_no_access(self):
        self.assertFalse(can_edit_claude_group(self.plain_reader, self.diabet))
        self.assertFalse(can_edit_claude_ration(self.plain_reader, self.diabet_ration))

    def test_editor_edits_everything_including_limited_editors_work(self):
        for group in (self.diabet, self.foreign, self.own):
            self.assertTrue(can_edit_claude_group(self.editor, group))
        for ration in (self.own_ration, self.foreign_ration, self.diabet_ration):
            self.assertTrue(can_edit_claude_ration(self.editor, ration))

    # ── запросы ──────────────────────────────────────────────────────────────

    def test_limited_editor_cannot_touch_foreign_ration(self):
        self.client.force_login(self.limited)
        resp = self.client.post(
            reverse("claude_ration_delete", args=[self.foreign_ration.pk])
        )
        self.assertEqual(resp.status_code, 403)
        self.assertTrue(ClaudeRation.objects.filter(pk=self.foreign_ration.pk).exists())

    def test_limited_editor_deletes_own_ration(self):
        self.client.force_login(self.limited)
        resp = self.client.post(
            reverse("claude_ration_delete", args=[self.own_ration.pk])
        )
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(ClaudeRation.objects.filter(pk=self.own_ration.pk).exists())

    def test_limited_editor_edits_diabet_group(self):
        self.client.force_login(self.limited)
        resp = self.client.post(
            reverse("claude_group_edit", args=[self.diabet.pk]),
            {"name": "ДИАБЕТ 1800 (правка)", "description": ""},
        )
        self.assertEqual(resp.status_code, 302)
        self.diabet.refresh_from_db()
        self.assertEqual(self.diabet.name, "ДИАБЕТ 1800 (правка)")

    def test_limited_editor_cannot_edit_foreign_group(self):
        self.client.force_login(self.limited)
        resp = self.client.post(
            reverse("claude_group_edit", args=[self.foreign.pk]),
            {"name": "Переименовал чужое", "description": ""},
        )
        self.assertEqual(resp.status_code, 403)
        self.foreign.refresh_from_db()
        self.assertEqual(self.foreign.name, "Рационы 1200 (тест)")

    def test_created_group_belongs_to_author(self):
        self.client.force_login(self.limited)
        self.client.post(reverse("claude_group_create"), {"name": "Новая", "description": ""})
        created = ClaudeRationGroup.objects.get(name="Новая")
        self.assertEqual(created.created_by, self.limited)
        self.assertTrue(can_edit_claude_group(self.limited, created))

    def test_plain_reader_blocked_by_middleware(self):
        self.client.force_login(self.plain_reader)
        resp = self.client.post(reverse("claude_group_create"), {"name": "Нельзя"})
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(ClaudeRationGroup.objects.filter(name="Нельзя").exists())

    def test_limited_editor_still_blocked_outside_claude(self):
        """Право открывает только вкладку Claude — калоражи остаются закрыты."""
        self.client.force_login(self.limited)
        resp = self.client.post(reverse("calorie_create"), {"name": "1300", "kcal": 1300})
        self.assertEqual(resp.status_code, 403)

    def test_limited_editor_cannot_move_foreign_ration(self):
        self.client.force_login(self.limited)
        resp = self.client.post(
            reverse("claude_ration_reorder"),
            data={
                "ration_id": self.foreign_ration.pk,
                "target_group_id": self.own.pk,
                "ordered_ids": [self.foreign_ration.pk],
            },
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 403)
        self.foreign_ration.refresh_from_db()
        self.assertEqual(self.foreign_ration.group_id, self.foreign.pk)


class JktCategoryTests(TestCase):
    """Категория «ЖКТ» и её сопоставление с группой внешнего меню iiko."""

    IIKO_GROUP = "ЖКТ (желудочно кишечный тракт)"

    def test_category_seeded_by_migration(self):
        from .models import IikoCategoryMap, MealCategory

        category = MealCategory.objects.get(key="jkt")
        self.assertEqual(str(category), "ЖКТ")
        self.assertEqual(
            IikoCategoryMap.objects.get(iiko_name=self.IIKO_GROUP).slot_key, "jkt"
        )

    def test_sync_maps_iiko_group_to_jkt(self):
        from .iiko_sync import _build_label_to_key, _pick_category

        cat_display, slot_key = _pick_category(
            {"category_names": [self.IIKO_GROUP]}, _build_label_to_key()
        )
        self.assertEqual(slot_key, "jkt")
        self.assertEqual(cat_display, self.IIKO_GROUP)

    def test_catalog_filter_finds_jkt_dishes(self):
        from .models import MealCategory

        user = User.objects.create_user("adil", password="x")
        dish = Product.objects.create(name="ЖКТ* Упак Винегрет (1порц)", article="11237")
        dish.meal_categories.set([MealCategory.objects.get(key="jkt")])
        Product.objects.create(name="ПП* Упак Бигус (1порц)", article="00381")

        self.client.force_login(user)
        html = self.client.get(reverse("product_list"), {"meal_category": "jkt"}).content.decode()
        self.assertIn("Винегрет", html)
        self.assertNotIn("Бигус", html)


class OrderSheetAppendixTests(TestCase):
    """Шапка «Приложение к договору поставки» в заявочном листе."""

    def _sheet(self, with_price):
        import io
        from openpyxl import load_workbook
        from .views import _build_order_sheet_xlsx

        product = Product.objects.create(
            name="ПП* Упак Блины с мясом (3шт)", article="29272",
            packing="порц", sale_price=Decimal("921"),
        )
        resp = _build_order_sheet_xlsx([product], with_price=with_price)
        return load_workbook(io.BytesIO(resp.content)).active

    def test_without_price_is_appendix_1(self):
        ws = self._sheet(with_price=False)
        self.assertEqual(ws["C1"].value, "ПРИЛОЖЕНИЕ № 1")
        self.assertTrue(ws["C1"].font.bold)
        self.assertEqual(ws["C2"].value, "к Договору поставки № ______")
        self.assertEqual(ws["C3"].value, "от «____» __________ 20___ г.")
        self.assertEqual(ws["C1"].alignment.horizontal, "right")
        # тянется до последней колонки «ИТОГО»
        self.assertIn("C1:D1", [str(r) for r in ws.merged_cells.ranges])

    def test_with_price_is_appendix_2(self):
        ws = self._sheet(with_price=True)
        self.assertEqual(ws["C1"].value, "ПРИЛОЖЕНИЕ № 2")
        self.assertIn("C1:E1", [str(r) for r in ws.merged_cells.ranges])

    def test_existing_header_stays_in_place(self):
        """Шапка приложения не сдвигает строки ручного файла."""
        ws = self._sheet(with_price=True)
        self.assertEqual(ws["A1"].value, 1)
        self.assertEqual(ws["B1"].value, 'Заявочный лист "Фуд завод"')
        self.assertEqual(ws["B4"].value, "O-Live")
        self.assertEqual(ws["A8"].value, "Артикул")
        self.assertEqual(ws["D8"].value, "Цена прод.")
        self.assertEqual(ws["A11"].value, "29272")


class ClaudeCatalogTests(TestCase):
    """Каталог, который уходит в Claude при сборке рациона."""

    def test_catalog_carries_composition(self):
        """Без состава нельзя выполнить «без сахара»: по названию не видно."""
        from .claude_rations import _build_catalog

        # composition_clean не задаём: он пересчитывается из сырого состава
        # в Product.save() по справочнику ингредиентов
        product = Product.objects.create(
            name="ПП* Упак Поке боул с уткой (1порц)",
            composition="Киноа, утиное филе, сахар, соус ореховый",
            kcal_per_serving=Decimal("520"),
        )
        row = _build_catalog([product])[0]
        self.assertIn("ахар", row["composition"])     # «Сахар» после очистки
        self.assertEqual(row["id"], product.pk)

    def test_catalog_falls_back_to_raw_composition(self):
        from .claude_rations import _build_catalog

        product = Product.objects.create(
            name="Блюдо без чистого состава", composition="Сырой состав из iiko",
        )
        self.assertEqual(_build_catalog([product])[0]["composition"], "Сырой состав из iiko")


class ProductCardTests(TestCase):
    """Карточка блюда, которую открывает клик по блюду в слоте рациона."""

    def setUp(self):
        self.user = User.objects.create_user("adil", password="x")
        self.product = Product.objects.create(
            name="ПП* Упак Блины с мясом (3шт)", article="29272",
            iiko_category="ПП* ЗАВТРАК", net_weight=Decimal("0.265"),
            cost=Decimal("551"), sale_price=Decimal("921"),
            composition="Мука, молоко, фарш говяжий, яйцо, соль",
            protein_per_serving=Decimal("25.1"), fat_per_serving=Decimal("28.2"),
            carbs_per_serving=Decimal("35.0"), kcal_per_serving=Decimal("494"),
            protein=Decimal("9.5"), fat=Decimal("10.6"),
            carbs=Decimal("13.2"), kcal_per_100=Decimal("186"),
        )

    def test_card_returns_full_info(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse("product_card", args=[self.product.pk]))
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["name"], "ПП* Упак Блины с мясом (3шт)")
        self.assertEqual(data["article"], "29272")
        self.assertIn("фарш говяжий", data["composition"])
        self.assertNotIn("allergens", data)      # аллергены с платформы убраны
        self.assertEqual(data["weight_g"], 265.0)      # кг из базы -> граммы
        self.assertEqual(data["sale_price"], 921.0)
        self.assertEqual(data["per_serving"]["kcal"], 494.0)
        self.assertEqual(data["per_100"]["kcal"], 186.0)

    def test_card_needs_login(self):
        resp = self.client.get(reverse("product_card", args=[self.product.pk]))
        self.assertEqual(resp.status_code, 302)        # уводит на /login/

    def test_ration_editor_makes_dish_clickable(self):
        group = ClaudeRationGroup.objects.create(name="ДИАБЕТ 2200")
        ration = ClaudeRation.objects.create(
            group=group, name="Диабет 2200 день 7", kcal_category=2200,
        )
        ClaudeRationSlot.objects.create(
            ration=ration, product=self.product, slot_type="breakfast",
            meal_time=MealTime.objects.create(name="Завтрак-1"),
        )
        self.client.force_login(self.user)
        html = self.client.get(
            reverse("claude_ration_edit", args=[ration.pk])
        ).content.decode()
        self.assertIn(f"openDishModal({self.product.pk})", html)
